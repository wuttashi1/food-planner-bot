import io
import math
import zipfile
import pytest
from openpyxl import load_workbook
from app.services.domain import canonical, nutrition, quantity, MenuData
from app.services.excel import export_workbook, parse_workbook
from app.seed.demo import demo


def mutate(fn):
    wb = load_workbook(io.BytesIO(export_workbook(demo())))
    fn(wb)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_units():
    assert canonical(1.2, "kg") == (1200, "g")
    assert canonical(1.5, "l") == (1500, "ml")
    assert canonical(2, "pcs") == (2, "pcs")
    assert quantity(1200, "g") == "1.2 кг"
    with pytest.raises(ValueError):
        canonical(1, "cups")


def test_nutrition_servings_and_overrides():
    r = demo().recipes[2]
    total = nutrition(r)
    assert total["kcal"] == 664.5
    r.servings = 2
    assert nutrition(r)["kcal"] == 332.25
    assert nutrition(r, 2) == total
    assert nutrition(r, 2, {"chicken": 220})["kcal"] == 708.5


def test_seed_week_kcal():
    data = demo()
    recipes = {r.recipe_id: r for r in data.recipes}
    for day in {m.day for m in data.meals}:
        kcal = sum(nutrition(recipes[m.recipe_id], m.portion)["kcal"] for m in data.meals if m.day == day)
        assert 2500 <= kcal <= 2600, (day, kcal)


def test_template_structure():
    raw = export_workbook()
    wb = load_workbook(io.BytesIO(raw))
    assert wb.sheetnames == ["README", "MENU", "RECIPES", "INGREDIENTS", "SETTINGS", "NOTIFICATIONS"]
    assert wb["MENU"].freeze_panes == "A2"
    assert wb["MENU"]["A1"].value == "menu_id"
    with pytest.raises(ValueError, match="MENU пуст"):
        parse_workbook(raw)


def test_export_import_roundtrip():
    data = demo()
    data.settings = {"timezone": "Europe/Berlin", "disabled_meals": ["snack"], "notifications_enabled": False}
    result, warnings = parse_workbook(export_workbook(data))
    assert not warnings
    assert result.model_dump() == data.model_dump()


def test_formula_injection_export_text():
    data = demo()
    data.recipes[0].recipe_name = '=HYPERLINK("https://evil.invalid")'
    data.meals[0].meal_name = "+1+1"
    wb = load_workbook(io.BytesIO(export_workbook(data)))
    assert wb["RECIPES"]["B2"].data_type == "s"
    assert wb["RECIPES"]["B2"].quotePrefix
    result, _ = parse_workbook(export_workbook(data))
    assert result.recipes[0].recipe_name.startswith("=")


@pytest.mark.parametrize(
    "fn",
    [
        lambda w: w.remove(w["RECIPES"]),
        lambda w: setattr(w["MENU"]["G2"], "value", "missing"),
        lambda w: setattr(w["MENU"]["C2"], "value", "funday"),
        lambda w: setattr(w["INGREDIENTS"]["D2"], "value", -1),
        lambda w: setattr(w["INGREDIENTS"]["E2"], "value", "cups"),
        lambda w: setattr(w["MENU"]["I2"], "value", "=1+2"),
        lambda w: setattr(w["MENU"]["A3"], "value", 99),
        lambda w: w["RECIPES"].append([c.value for c in w["RECIPES"][2]]),
        lambda w: w["INGREDIENTS"].append([c.value for c in w["INGREDIENTS"][2]]),
        lambda w: setattr(w["MENU"]["I2"], "value", "nan"),
        lambda w: setattr(w["MENU"]["I2"], "value", -10),
    ],
)
def test_validation_rejects_bad_files(fn):
    with pytest.raises(ValueError):
        parse_workbook(mutate(fn))


def test_nutrition_mismatch_warns_and_recalculates():
    raw = mutate(lambda w: setattr(w["MENU"]["I2"], "value", 9999))
    result, warnings = parse_workbook(raw)
    assert warnings
    assert nutrition(result.recipes[0])["kcal"] != 9999


def test_reject_size_and_macro():
    with pytest.raises(ValueError):
        parse_workbook(export_workbook(demo()), max_bytes=10)
    raw = io.BytesIO(export_workbook(demo()))
    with zipfile.ZipFile(raw, "a") as z:
        z.writestr("xl/vbaProject.bin", b"bad")
    with pytest.raises(ValueError, match="Макросы"):
        parse_workbook(raw.getvalue())


def test_reject_nonfinite_product():
    data = demo().model_dump()
    data["recipes"][0]["ingredients"][0]["amount"] = math.inf
    with pytest.raises(ValueError):
        MenuData.model_validate(data)


@pytest.mark.parametrize(
    "settings",
    [
        {"gpt_prompt": 123},
        {"name": True},
        {"disabled_meals": [{}]},
        {"timezone": "Invalid/Timezone"},
        {"target_weight": float("nan")},
    ],
)
def test_reject_bad_settings(settings):
    data = demo().model_dump()
    data["settings"] = settings
    with pytest.raises(ValueError):
        MenuData.model_validate(data)


def test_fractional_menu_id_rejected():
    with pytest.raises(ValueError, match="целым числом"):
        parse_workbook(mutate(lambda w: setattr(w["MENU"]["A2"], "value", 1.5)))


def test_excel_native_time_cells():
    from datetime import time
    from app.services.domain import NoticeData

    data = demo()
    data.notifications = [NoticeData(notification_type="meal", meal_type="lunch", time="13:00")]
    wb = load_workbook(io.BytesIO(export_workbook(data)))
    wb["NOTIFICATIONS"]["D2"] = time(13, 0)
    out = io.BytesIO()
    wb.save(out)
    result, _ = parse_workbook(out.getvalue())
    assert result.notifications[0].time == "13:00"
