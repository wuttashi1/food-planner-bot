"""Strict, formula-free workbook interchange. No writes during validation."""

import io
from datetime import time
import json
import zipfile
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from app.services.domain import MenuData, RecipeData, MealData, Product, NoticeData, DAYS, nutrition

COLUMNS = {
    "MENU": "menu_id menu_version day meal_type meal_order meal_name recipe_id portion kcal protein fat carbs enabled".split(),
    "RECIPES": "recipe_id recipe_name servings prep_time cook_time instructions".split(),
    "INGREDIENTS": "recipe_id ingredient_id ingredient_name amount unit category kcal_per_100 protein_per_100 fat_per_100 carbs_per_100".split(),
    "SETTINGS": ["key", "value"],
    "NOTIFICATIONS": "notification_type meal_type enabled time offset_minutes".split(),
}
GPT_PROMPT = """Я прикрепил Excel-файл из Telegram Food Planner.
Измени рацион согласно задаче: [ОПИШИТЕ ЗАДАЧУ].
Верни .xlsx. Сохрани названия листов и обязательные колонки.
Не удаляй recipe_id. Не создавай дубли ingredient_id без необходимости.
MENU должен ссылаться на существующие RECIPES. INGREDIENTS должен содержать все продукты рецептов.
Обнови рецепты и ингредиенты при замене блюда. Пересчитай kcal, protein, fat, carbs из ингредиентов с учётом servings и portion.
Пищевая ценность указана на 100 g/ml/pcs canonical unit, даже если количество указано в kg/l.
Не используй формулы и макросы. Заполни запрошенные дни. Проверь целостность данных.
Сохрани совместимость с Food Planner Telegram Bot."""
README = [
    "FOOD PLANNER TEMPLATE",
    "Не переименовывайте системные листы и обязательные колонки.",
    "MENU — повторяющийся недельный рацион. day: monday … sunday.",
    "RECIPES — рецепты. servings — число порций в полном рецепте.",
    "INGREDIENTS — количества на полный рецепт. MENU.portion — число порций участника.",
    "Единицы: g, kg, ml, l, pcs. БЖУ всегда на 100 g/ml/pcs, не на kg/l.",
    "Одинаковый ingredient_id означает один продукт с одинаковыми единицами и БЖУ.",
    "kcal, protein, fat, carbs в MENU справочные: бот пересчитывает их из ингредиентов.",
    "SETTINGS и NOTIFICATIONS применяются только к пользователю, который импортирует файл.",
    "Дополнительная колонка weekday в NOTIFICATIONS: 0=понедельник … 6=воскресенье.",
    "Чистый шаблон не содержит меню. Заполните MENU, RECIPES и INGREDIENTS.",
    "Формулы и макросы запрещены. Сохраните .xlsx и отправьте через /import.",
    "ChatGPT: /export → загрузите файл в ChatGPT → скопируйте промт ниже.",
    "Скачайте изменённый .xlsx → /import → проверьте preview → подтвердите.",
    "Личные порции и отметки покупок в файл общего меню не входят.",
    GPT_PROMPT,
]


def write_cell(cell, value):
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    cell.value = value
    if isinstance(value, str):
        # Explicit text cell prevents formula injection without altering identifiers.
        cell.data_type = "s"
        if value.startswith(("=", "+", "-", "@")):
            cell.quotePrefix = True


def export_workbook(data=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    ws.column_dimensions["A"].width = 110
    for i, text in enumerate(README, 1):
        write_cell(ws.cell(i, 1), text)
        ws.cell(i, 1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = 220 if i == len(README) else 44
    ws["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="173C4A")
    for name, columns in COLUMNS.items():
        sheet = wb.create_sheet(name)
        sheet.append(columns + (["weekday"] if name == "NOTIFICATIONS" else []))
    if data:
        recipes = {r.recipe_id: r for r in data.recipes}
        for m in data.meals:
            values = {
                "menu_id": data.menu_id,
                "menu_version": data.menu_version,
                **m.model_dump(),
                **nutrition(recipes[m.recipe_id], m.portion),
            }
            wb["MENU"].append([values[k] for k in COLUMNS["MENU"]])
        for r in data.recipes:
            wb["RECIPES"].append([getattr(r, k) for k in COLUMNS["RECIPES"]])
            for p in r.ingredients:
                values = {"recipe_id": r.recipe_id, **p.model_dump()}
                wb["INGREDIENTS"].append([values[k] for k in COLUMNS["INGREDIENTS"]])
        for k, value in data.settings.items():
            wb["SETTINGS"].append([k, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value])
        for n in data.notifications:
            wb["NOTIFICATIONS"].append([getattr(n, k) for k in COLUMNS["NOTIFICATIONS"] + ["weekday"]])
    else:
        wb["SETTINGS"].append(["timezone", "Europe/Berlin"])
    for sheet in list(wb)[1:]:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        for row in sheet:
            for cell in row:
                write_cell(cell, cell.value)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if cell.row == 1:
                    cell.fill = PatternFill("solid", fgColor="173C4A")
                    cell.font = Font(bold=True, color="FFFFFF")
                elif cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="EFF6F7")
                if isinstance(cell.value, float):
                    cell.number_format = "0.00"
        for cell in sheet[1]:
            sheet.column_dimensions[cell.column_letter].width = (
                52 if cell.value in ("instructions", "value") else 28 if "name" in str(cell.value) else max(18, len(str(cell.value)) + 2)
            )
        sheet.row_dimensions[1].height = 32
        for i in range(2, sheet.max_row + 1):
            sheet.row_dimensions[i].height = 70 if sheet.title == "RECIPES" else 32
    dv = DataValidation(type="list", formula1='"' + ",".join(DAYS) + '"')
    wb["MENU"].add_data_validation(dv)
    dv.add("C2:C701")
    dv = DataValidation(type="list", formula1='"g,kg,ml,l,pcs"')
    wb["INGREDIENTS"].add_data_validation(dv)
    dv.add("E2:E10001")
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def parse_workbook(raw, max_bytes=5 * 1024 * 1024):
    if len(raw) > max_bytes:
        raise ValueError("Файл слишком большой.")
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            entries = archive.infolist()
            if len(entries) > 1000 or sum(e.file_size for e in entries) > 30 * 1024 * 1024:
                raise ValueError("Слишком большой распакованный Excel.")
            if any("vbaproject" in e.filename.lower() or "externallinks" in e.filename.lower() for e in entries):
                raise ValueError("Макросы и внешние ссылки запрещены.")
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    except (zipfile.BadZipFile, KeyError, OSError) as exc:
        raise ValueError("Не удалось открыть .xlsx.") from exc
    try:
        missing = {"README", *COLUMNS} - set(wb.sheetnames)
        if missing:
            raise ValueError("Нет листов: " + ", ".join(sorted(missing)))
        tables = {}
        for sheet in wb:
            # Ignore forged worksheet dimensions; bound actual rows/cells.
            sheet.reset_dimensions()
            rows = sheet.iter_rows()
            header = next(rows, ())
            if len(header) > 64:
                raise ValueError("Слишком много колонок.")
            columns = [c.value for c in header]
            if sheet.title in COLUMNS:
                if len(set(columns)) != len(columns) or set(COLUMNS[sheet.title]) - set(columns):
                    raise ValueError(f"{sheet.title}: нет обязательных колонок или есть дубли.")
            table = []
            for number, row in enumerate(rows, 2):
                if number > 10002 or len(row) > 64:
                    raise ValueError(f"{sheet.title}: превышен лимит строк/колонок.")
                for c in row:
                    if c.data_type == "f":
                        raise ValueError(f"{sheet.title}!{c.coordinate}: формулы запрещены.")
                    if isinstance(c.value, str) and len(c.value) > 12000:
                        raise ValueError("Слишком длинный текст в ячейке.")
                if any(c.value is not None for c in row):
                    table.append(dict(zip(columns, [c.value for c in row])))
            for c in header:
                if c.data_type == "f":
                    raise ValueError("Формулы запрещены.")
            tables[sheet.title] = table
        if not tables["MENU"]:
            raise ValueError("MENU пуст. Заполните шаблон перед импортом.")
        ingredients = {}
        for row in tables["INGREDIENTS"]:
            rid = str(row["recipe_id"])
            values = {k: row[k] for k in COLUMNS["INGREDIENTS"] if k != "recipe_id"}
            values["ingredient_id"] = str(values["ingredient_id"])
            ingredients.setdefault(rid, []).append(Product(**values))
        recipes = []
        for row in tables["RECIPES"]:
            rid = str(row["recipe_id"])
            values = {k: row[k] for k in COLUMNS["RECIPES"]}
            values["recipe_id"] = rid
            values["instructions"] = values["instructions"] or ""
            recipes.append(RecipeData(**values, ingredients=ingredients.get(rid, [])))
        if set(ingredients) - {r.recipe_id for r in recipes}:
            raise ValueError("INGREDIENTS ссылается на отсутствующий рецепт.")
        meals, ids, versions = [], set(), set()
        for row in tables["MENU"]:
            for key in ("menu_id", "menu_version"):
                value = row[key]
                if isinstance(value, bool) or value is None or float(value) != int(value):
                    raise ValueError(f"{key} должен быть целым числом.")
            ids.add(int(row["menu_id"]))
            versions.add(int(row["menu_version"]))
            values = {k: row[k] for k in MealData.model_fields}
            values["recipe_id"] = str(values["recipe_id"])
            meals.append(MealData(**values))
        if len(ids) != 1 or len(versions) != 1:
            raise ValueError("В одном файле должно быть одно menu_id и menu_version.")
        settings = {}
        for row in tables["SETTINGS"]:
            key, value = row["key"], row["value"]
            if key in settings:
                raise ValueError("Повтор ключа SETTINGS.")
            if key == "disabled_meals" and isinstance(value, str):
                value = json.loads(value)
            settings[key] = value
        notices = []
        for row in tables["NOTIFICATIONS"]:
            values = {k: row[k] for k in COLUMNS["NOTIFICATIONS"]}
            values["meal_type"] = values["meal_type"] or ""
            if isinstance(values["time"], time):
                if values["time"].second or values["time"].microsecond:
                    raise ValueError("Уведомления поддерживают точность до минуты.")
                values["time"] = values["time"].strftime("%H:%M")
            values["weekday"] = row.get("weekday") if row.get("weekday") is not None else 5
            notices.append(NoticeData(**values))
        data = MenuData(
            menu_id=ids.pop(), menu_version=versions.pop(), meals=meals, recipes=recipes, settings=settings, notifications=notices
        )
        warnings = []
        recs = {r.recipe_id: r for r in recipes}
        for i, (m, row) in enumerate(zip(meals, tables["MENU"]), 2):
            calculated = nutrition(recs[m.recipe_id], m.portion)
            for k in ("kcal", "protein", "fat", "carbs"):
                value = float(row[k])
                if not 0 <= value <= 100000:
                    raise ValueError(f"MENU!{i}: некорректное {k}.")
                if abs(value - calculated[k]) > max(1, calculated[k] * 0.02):
                    warnings.append(f"MENU строка {i}: {k} пересчитано {value:g} → {calculated[k]:g}")
        if len({m.day for m in meals}) < 7:
            warnings.append("Заполнены не все 7 дней.")
        return data, warnings
    finally:
        wb.close()


def preview(data):
    recs = {r.recipe_id: r for r in data.recipes}
    kcal = sum(nutrition(recs[m.recipe_id], m.portion)["kcal"] for m in data.meals if m.enabled)
    days = len({m.day for m in data.meals})
    return f"📅 Дней: {days}\n🍽 Приёмов: {len(data.meals)}\n🥘 Рецептов: {len(data.recipes)}\n🥕 Продуктов: {len({p.ingredient_id for r in data.recipes for p in r.ingredients})}\n🔥 Среднее: {kcal / days:.0f} kcal/день"
